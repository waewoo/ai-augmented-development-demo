"""Build the editable 16:9 presentation and speaker notes.

The deck intentionally uses native PowerPoint shapes and text instead of
screenshots so that the presenter can adapt it without a design tool.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "presentation" / "premiers-pas-developpement-augmente-ia.pptx"

W = Inches(13.333)
H = Inches(7.5)
BG = RGBColor(247, 249, 252)
INK = RGBColor(22, 31, 49)
MUTED = RGBColor(91, 104, 124)
NAVY = RGBColor(18, 34, 60)
BLUE = RGBColor(61, 128, 224)
CYAN = RGBColor(46, 184, 184)
ORANGE = RGBColor(242, 153, 74)
RED = RGBColor(207, 79, 79)
WHITE = RGBColor(255, 255, 255)
PALE_BLUE = RGBColor(229, 239, 253)
PALE_CYAN = RGBColor(226, 247, 245)
PALE_ORANGE = RGBColor(255, 241, 224)
PALE_RED = RGBColor(252, 233, 233)


def set_bg(slide: object, color: RGBColor = BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def shape(
    slide: object,
    kind: MSO_SHAPE,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor = WHITE,
    line: RGBColor | None = None,
    radius: bool = False,
):
    chosen = MSO_SHAPE.ROUNDED_RECTANGLE if radius else kind
    item = slide.shapes.add_shape(chosen, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = fill
    item.line.color.rgb = line or fill
    return item


def text(
    slide: object,
    value: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 18,
    color: RGBColor = INK,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    font: str = "Aptos",
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def bullet_list(
    slide: object,
    items: list[str],
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 18,
    color: RGBColor = INK,
    bullet_color: RGBColor = BLUE,
):
    for index, item in enumerate(items):
        yy = y + index * (h / max(len(items), 1))
        shape(slide, MSO_SHAPE.OVAL, x, yy + 0.1, 0.12, 0.12, bullet_color)
        text(slide, item, x + 0.25, yy, w - 0.25, h / max(len(items), 1), size, color)


def header(slide: object, number: int, title: str, kicker: str, duration: str):
    text(slide, kicker.upper(), 0.65, 0.36, 5.8, 0.25, 10, BLUE, True)
    text(slide, title, 0.65, 0.66, 11.6, 0.55, 27, INK, True)
    shape(slide, MSO_SHAPE.RECTANGLE, 0.65, 1.38, 12.05, 0.025, PALE_BLUE)
    text(
        slide,
        f"{number:02d}  ·  {duration}",
        11.25,
        7.08,
        1.4,
        0.18,
        9,
        MUTED,
        False,
        PP_ALIGN.RIGHT,
    )


def pill(
    slide: object,
    label: str,
    x: float,
    y: float,
    w: float,
    fill: RGBColor,
    color: RGBColor = INK,
):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.38, fill, fill, True)
    text(
        slide,
        label,
        x + 0.08,
        y + 0.07,
        w - 0.16,
        0.22,
        11,
        color,
        True,
        PP_ALIGN.CENTER,
    )


def card(
    slide: object,
    title: str,
    body: str,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor = WHITE,
    accent: RGBColor = BLUE,
    title_size: float = 16,
    body_size: float = 13,
):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, fill, True)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.08, h, accent)
    text(slide, title, x + 0.25, y + 0.2, w - 0.45, 0.3, title_size, INK, True)
    text(slide, body, x + 0.25, y + 0.64, w - 0.45, h - 0.78, body_size, MUTED)


def arrow(slide: object, x: float, y: float, w: float, color: RGBColor = BLUE):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.24)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.color.rgb = color


def add_notes(slide: object, notes: str):
    frame = slide.notes_slide.notes_text_frame
    frame.clear()
    frame.text = notes


def make_deck() -> None:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    def new_slide(number: int, title: str, kicker: str, duration: str):
        slide = prs.slides.add_slide(blank)
        set_bg(slide)
        header(slide, number, title, kicker, duration)
        return slide

    notes = []

    # 1
    slide = prs.slides.add_slide(blank)
    set_bg(slide, NAVY)
    shape(slide, MSO_SHAPE.OVAL, 10.1, -0.6, 4.2, 4.2, BLUE)
    shape(slide, MSO_SHAPE.OVAL, 11.25, 4.55, 2.9, 2.9, CYAN)
    pill(
        slide,
        "UNE MÉTHODE · DES GARDE-FOUS · UN EXEMPLE",
        0.75,
        0.72,
        4.7,
        PALE_BLUE,
        NAVY,
    )
    text(
        slide,
        "Premiers pas vers\nle développement\naugmenté par l’IA",
        0.75,
        1.55,
        8.5,
        2.2,
        34,
        WHITE,
        True,
    )
    text(
        slide,
        "Une heure pour repartir avec une méthode reproductible — pas pour devenir expert.",
        0.78,
        4.22,
        7.3,
        0.55,
        18,
        RGBColor(212, 224, 242),
    )
    text(
        slide,
        "Kilo Code comme outil de démonstration · principes transposables",
        0.78,
        6.55,
        7.5,
        0.25,
        11,
        RGBColor(180, 198, 221),
    )
    text(
        slide,
        "01  ·  1 min",
        11.4,
        7.08,
        1.2,
        0.18,
        9,
        RGBColor(180, 198, 221),
        False,
        PP_ALIGN.RIGHT,
    )
    notes.append(
        """Durée : 1 min. Dire : « En une heure, vous ne deviendrez pas développeur IA. Vous repartirez toutefois avec une méthode, des garde-fous et un exemple que vous pourrez reproduire. » Annoncer le fil : contexte, rules, skills, workflow, frameworks, contrôles. Kilo Code est l’outil visible de la démonstration, pas une recommandation exclusive. Transition : Pour comprendre pourquoi cette méthode existe, regardons ce qui a changé."""
    )

    # 2
    slide = new_slide(2, "Pourquoi maintenant ?", "Contexte", "2 min")
    card(
        slide,
        "Avant",
        "Complétion locale\nUne ligne, une fonction, une réponse.",
        0.75,
        2.0,
        3.55,
        2.1,
        WHITE,
        MUTED,
    )
    arrow(slide, 4.65, 2.78, 0.72, ORANGE)
    card(
        slide,
        "Maintenant",
        "Explorer · planifier · modifier\n· tester un projet.",
        5.55,
        2.0,
        3.55,
        2.1,
        PALE_BLUE,
        BLUE,
    )
    arrow(slide, 9.45, 2.78, 0.72, ORANGE)
    card(
        slide,
        "Responsabilité",
        "L’agent propose et exécute.\nLe développeur décide et vérifie.",
        10.35,
        2.0,
        2.2,
        2.1,
        PALE_ORANGE,
        ORANGE,
        14,
        12,
    )
    text(
        slide,
        "Plus de portée = plus de valeur potentielle… et plus de surface d’erreur.",
        1.0,
        5.2,
        10.7,
        0.55,
        22,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    notes.append(
        """Durée : 2 min. La complétion aide sur une ligne ou une fonction. Un agent peut explorer, planifier, appeler des outils, modifier plusieurs fichiers et vérifier. Cela augmente aussi la portée des erreurs. Le développeur reste responsable des décisions et du résultat. Transition : La différence ne tient donc pas seulement au modèle ; elle tient à l’environnement de travail."""
    )

    # 3
    slide = new_slide(3, "Du modèle à l’agent", "Vocabulaire", "3 min")
    card(
        slide,
        "Modèle",
        "Produit une réponse à partir des informations reçues.",
        0.8,
        2.0,
        3.1,
        2.25,
        WHITE,
        MUTED,
    )
    arrow(slide, 4.2, 2.95, 0.72, BLUE)
    card(
        slide,
        "Agent",
        "Observe\nPlanifie\nUtilise des outils\nModifie\nVérifie",
        5.1,
        1.75,
        3.0,
        2.75,
        PALE_BLUE,
        BLUE,
    )
    arrow(slide, 8.35, 2.95, 0.72, BLUE)
    card(
        slide,
        "Environnement",
        "Fichiers · terminal · Git\nTests · recherche · MCP\nPermissions · hooks",
        9.25,
        2.0,
        3.2,
        2.25,
        PALE_CYAN,
        CYAN,
    )
    pill(
        slide,
        "MCP = exposer outils / sources de façon standardisée",
        2.0,
        5.35,
        4.3,
        PALE_ORANGE,
        INK,
    )
    pill(
        slide,
        "Hooks = action avant / après un événement",
        6.55,
        5.35,
        3.8,
        PALE_ORANGE,
        INK,
    )
    pill(
        slide,
        "Sous-agents = délégation + coût + coordination",
        10.55,
        5.35,
        2.0,
        PALE_RED,
        INK,
    )
    notes.append(
        """Durée : 3 min. Expliquer simplement modèle versus agent. Poser le vocabulaire : outils de lecture/écriture, terminal, Git, tests, recherche ; MCP comme protocole standardisé ; hooks comme actions avant ou après un événement ; sous-agents comme délégation spécialisée avec coût, coordination et complexité. Dire qu’un agent n’est ni autonome au sens absolu ni infaillible. Transition : Pour garder cette capacité sous contrôle, on rend les étapes visibles."""
    )

    # 4
    slide = new_slide(4, "Le workflow de référence", "Méthode", "2 min")
    steps = [
        "Comprendre",
        "Planifier",
        "Approuver",
        "Implémenter",
        "Vérifier",
        "Revoir le diff",
    ]
    for i, label in enumerate(steps):
        x = 0.75 + i * 2.05
        fill = PALE_BLUE if i in {1, 2} else WHITE
        card(
            slide,
            f"0{i + 1}",
            label,
            x,
            2.25,
            1.65,
            1.45,
            fill,
            BLUE if i in {1, 2} else MUTED,
            15,
            14,
        )
        if i < len(steps) - 1:
            arrow(slide, x + 1.7, 2.85, 0.27, ORANGE)
    text(
        slide,
        "Le plan est un point de contrôle humain, pas une formalité.",
        1.0,
        5.15,
        11.2,
        0.55,
        23,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    notes.append(
        """Durée : 2 min. Parcourir comprendre → planifier → faire approuver → implémenter → vérifier → revoir le diff. Insister sur l’approbation humaine : on peut réduire le périmètre ou corriger une mauvaise interprétation avant les modifications. Transition : Comprendre dépend de ce que l’agent peut réellement voir."""
    )

    # 5
    slide = new_slide(
        5, "Le contexte : premier facteur de qualité", "Contexte", "3 min"
    )
    items = [
        ("Code", BLUE),
        ("Demande + critères", CYAN),
        ("Architecture + conventions", ORANGE),
        ("Tests + commandes", RED),
        ("Sécurité + historique utile", MUTED),
    ]
    for i, (label, color) in enumerate(items):
        x = 0.9 + (i % 3) * 4.1
        y = 2.0 + (i // 3) * 1.45
        card(
            slide,
            label,
            "Contexte pertinent\nplutôt que tout charger",
            x,
            y,
            3.45,
            1.1,
            WHITE if i != 1 else PALE_BLUE,
            color,
            15,
            12,
        )
    text(
        slide,
        "Mauvais contexte → résultat plausible, souvent incorrect.",
        1.0,
        5.5,
        11.2,
        0.45,
        22,
        RED,
        True,
        PP_ALIGN.CENTER,
    )
    notes.append(
        """Durée : 3 min. Le contexte utile rassemble code, documentation, demande, critères, architecture, conventions, historique pertinent, commandes de validation et contraintes de sécurité. Insister : sélectionner le contexte pertinent plutôt que tout charger sans discernement. Phrase : « Je donne d’abord à l’agent les moyens de comprendre le projet. » Transition : Voyons cette découverte sans toucher à un fichier."""
    )

    # 6
    slide = new_slide(
        6, "Démonstration 1 — découvrir avant d’agir", "Démo · lecture seule", "3 min"
    )
    card(
        slide,
        "Demande",
        "Explore le dépôt\nsans modifier de fichier.",
        0.8,
        2.05,
        2.7,
        1.75,
        PALE_BLUE,
        BLUE,
    )
    arrow(slide, 3.7, 2.72, 0.7, BLUE)
    card(
        slide,
        "L’agent identifie",
        "Structure · point d’entrée\nTests · commandes\nRules · skills",
        4.6,
        1.75,
        3.25,
        2.35,
        WHITE,
        CYAN,
    )
    arrow(slide, 8.1, 2.72, 0.7, BLUE)
    card(
        slide,
        "Preuve",
        "Aucun fichier modifié\nLe terrain est compris\navant le code.",
        9.0,
        2.05,
        3.45,
        1.75,
        PALE_CYAN,
        CYAN,
    )
    pill(
        slide,
        "Prompt exact dans docs/demo-prompts.md",
        3.55,
        5.45,
        5.3,
        PALE_ORANGE,
        INK,
    )
    notes.append(
        """Durée : 3 min. Ouvrir le projet avec Kilo Code, rester en mode planification/lecture seule et saisir le prompt d’exploration. Montrer structure, app/main.py, tests, make check, kilo.jsonc, rules et skills. Ne pas lire ligne à ligne : relever deux observations. Solution de secours : demo-assets/expected-plan.md. Transition : Certaines attentes ne doivent pas être répétées à chaque demande : ce sont les rules."""
    )

    # 7
    slide = new_slide(
        7, "Les rules : rendre les attentes persistantes", "Rules", "2 min"
    )
    card(
        slide,
        "Une rule dit surtout…",
        "Ce qui doit être respecté\narchitecture · tests · sécurité\nformat de restitution",
        0.85,
        2.0,
        3.55,
        2.3,
        PALE_BLUE,
        BLUE,
    )
    arrow(slide, 4.65, 2.9, 0.75, ORANGE)
    card(
        slide,
        "Dans le projet",
        "kilo.jsonc\n→ instructions\n→ .kilo/rules/*.md",
        5.65,
        2.0,
        2.65,
        2.3,
        WHITE,
        CYAN,
    )
    arrow(slide, 8.6, 2.9, 0.75, ORANGE)
    card(
        slide,
        "Mais pas un contrôle absolu",
        "Best effort du modèle\nPermissions · CI · tests\nRevue humaine",
        9.6,
        2.0,
        2.75,
        2.3,
        PALE_ORANGE,
        ORANGE,
        14,
        12,
    )
    text(
        slide,
        "Lisible · auditable · persistant",
        2.15,
        5.45,
        9.0,
        0.4,
        22,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    notes.append(
        """Durée : 2 min. Une rule est une instruction persistante propre au projet ou à l’équipe. Elle dit principalement ce qui doit être respecté. Montrer kilo.jsonc et les quatre rules. Dire explicitement qu’une rule n’est pas une barrière technique : permissions, tests, CI et revue restent nécessaires. Transition : Mesurons l’effet avec deux plans, sans refaire la fonctionnalité."""
    )

    # 8
    slide = new_slide(8, "Sans rules / avec rules", "Comparaison", "2 min")
    card(
        slide,
        "Sans rules",
        "Ajouter status\nFiltrer\nAjouter quelques tests\n\n→ ambiguïtés non résolues",
        0.9,
        1.95,
        4.8,
        3.0,
        WHITE,
        MUTED,
        19,
        16,
    )
    card(
        slide,
        "Avec rules",
        "Fichiers ciblés\nType fermé\nTests : nominal + 422\nmake check + diff\n\n→ comportement reproductible",
        7.0,
        1.95,
        4.8,
        3.0,
        PALE_BLUE,
        BLUE,
        19,
        16,
    )
    text(
        slide,
        "On compare des plans, pas deux implémentations complètes.",
        1.1,
        5.65,
        10.9,
        0.42,
        19,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    notes.append(
        """Durée : 2 min. Afficher les deux fichiers Markdown préparés. Le plan sans rules est générique ; celui avec rules cible les fichiers, les tests et make check. Ne pas faire deux implémentations. Conclusion : les rules réduisent l’ambiguïté et rendent le comportement plus reproductible. Transition : Une rule dit surtout quoi respecter ; un skill décrit comment exécuter une tâche."""
    )

    # 9
    slide = new_slide(9, "Les skills : encapsuler une méthode", "Skills", "2 min")
    card(
        slide,
        "Rule",
        "Contraintes durables\n« respecte ceci »",
        0.9,
        2.05,
        3.25,
        2.1,
        PALE_BLUE,
        BLUE,
        19,
        16,
    )
    arrow(slide, 4.45, 2.8, 0.7, ORANGE)
    card(
        slide,
        "Skill",
        "Procédure réutilisable\n« fais cela ainsi »",
        5.25,
        2.05,
        3.25,
        2.1,
        PALE_CYAN,
        CYAN,
        19,
        16,
    )
    arrow(slide, 8.8, 2.8, 0.7, ORANGE)
    card(
        slide,
        "Résultat",
        "Méthode répétable\nplan · implement · review",
        9.6,
        2.05,
        2.75,
        2.1,
        PALE_ORANGE,
        ORANGE,
        16,
        14,
    )
    text(
        slide,
        ".kilo/skills/<nom>/SKILL.md",
        2.2,
        5.4,
        8.7,
        0.45,
        23,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    notes.append(
        """Durée : 2 min. Un skill est une procédure réutilisable : éléments à lire, étapes, contrôles et forme du résultat. Montrer le frontmatter name/description et les trois skills. La rule dit principalement quoi respecter ; le skill décrit comment exécuter une tâche ou un workflow. Transition : Utilisons plan-change sur une demande petite."""
    )

    # 10
    slide = new_slide(
        10, "Démonstration 2 — planifier avec un skill", "Démo · approbation", "4 min"
    )
    card(
        slide,
        "Demande",
        "Ajouter un filtre optionnel\n`status` sur `GET /tasks`",
        0.85,
        2.0,
        3.1,
        2.0,
        PALE_BLUE,
        BLUE,
        17,
        15,
    )
    arrow(slide, 4.15, 2.82, 0.7, BLUE)
    card(
        slide,
        "plan-change",
        "Lit code + rules + tests\nReformule les critères\nPropose fichiers + tests + risques",
        5.05,
        1.7,
        3.55,
        2.6,
        WHITE,
        CYAN,
        16,
        14,
    )
    arrow(slide, 8.8, 2.82, 0.7, BLUE)
    card(
        slide,
        "STOP",
        "Approbation humaine\navant toute écriture",
        9.7,
        2.0,
        2.65,
        2.0,
        PALE_ORANGE,
        ORANGE,
        19,
        15,
    )
    text(
        slide,
        "Le plan n’est pas une formalité : c’est le garde-fou le moins coûteux.",
        1.0,
        5.4,
        11.3,
        0.48,
        21,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    notes.append(
        """Durée : 4 min. Saisir le prompt plan-change. Laisser l’agent lire code, tests et rules. Attendre critères, fichiers, tests et risques. Arrêter avant toute modification. Poser la question : qu’est-ce qui ferait approuver ou refuser ce plan ? Marquer : « Le plan est un point de contrôle humain, pas une formalité. » Secours : expected-plan.md. Transition : Le plan est approuvé ; l’agent peut agir dans un périmètre connu."""
    )

    # 11
    slide = new_slide(
        11, "Démonstration 3 — implémenter et vérifier", "Démo · harnais", "5 min"
    )
    columns = [
        ("1", "Implémenter", PALE_BLUE, BLUE),
        ("2", "Tester", PALE_CYAN, CYAN),
        ("3", "Relire", PALE_ORANGE, ORANGE),
    ]
    for i, (num, title, fill, color) in enumerate(columns):
        x = 0.9 + i * 4.1
        card(
            slide,
            num + "  " + title,
            [
                "Filtre minimal\nPérimètre approuvé",
                "Sans filtre\nvalide · invalide 422",
                "Diff final\nrisques · scope drift",
            ][i],
            x,
            1.95,
            3.45,
            2.0,
            fill,
            color,
            18,
            15,
        )
        if i < 2:
            arrow(slide, x + 3.55, 2.8, 0.32, ORANGE)
    pill(slide, "make check", 5.0, 4.85, 3.3, NAVY, WHITE)
    text(
        slide,
        "Accepté parce que les contrôles sont verts — pas parce que l’agent dit « terminé ». ",
        1.05,
        5.58,
        11.1,
        0.5,
        19,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    notes.append(
        """Durée : 5 min. Saisir le prompt implement-change après approbation. Montrer les fichiers touchés, les tests et make check. Vérifier absence de filtre, filtre valide et 422. Afficher le diff. Dire : « Les tests et outils déterministes vérifient le résultat. » Puis : « Je regarde le diff et les contrôles, pas uniquement le résumé de l’agent. » Secours : expected-diff.md, expected-checks.md ou tag demo/implemented en worktree. Transition : Ce qui a protégé la modification forme le harnais."""
    )

    # 12
    slide = new_slide(
        12, "Le harnais : rendre l’agent vérifiable", "Contrôles", "3 min"
    )
    labels = [
        "Contexte",
        "Rules",
        "Outils",
        "Tests",
        "Lint / types",
        "CI",
        "Diff / permissions",
    ]
    for i, label in enumerate(labels):
        x = 0.75 + (i % 4) * 3.1
        y = 1.95 + (i // 4) * 1.45
        card(
            slide,
            str(i + 1),
            label,
            x,
            y,
            2.45,
            1.0,
            WHITE if i not in {3, 6} else PALE_BLUE,
            BLUE if i not in {3, 6} else CYAN,
            15,
            14,
        )
    text(
        slide,
        "Plus l’agent peut agir, plus le harnais doit être explicite.",
        1.0,
        5.45,
        11.25,
        0.5,
        22,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    notes.append(
        """Durée : 3 min. Définir : « Le harnais est l’ensemble du contexte, des règles, outils, tests, permissions et boucles de validation qui encadrent l’agent. » Montrer tests unitaires, intégration/API, lint, types, build/CI, revue du diff, permissions. Mentionner les hooks comme automatisations avant/après. Message : « Plus l’agent peut agir, plus le harnais doit être explicite. » Transition : Le changement respecte-t-il réellement la demande ?"""
    )

    # 13
    slide = new_slide(
        13, "Démonstration 4 — revoir sans modifier", "Démo · revue", "3 min"
    )
    card(
        slide,
        "Compare",
        "Demande initiale\nPlan approuvé\nRules",
        0.85,
        2.0,
        2.8,
        2.25,
        PALE_BLUE,
        BLUE,
        17,
        15,
    )
    arrow(slide, 3.85, 2.85, 0.7, BLUE)
    card(
        slide,
        "Inspecte",
        "Diff complet\nRésultats de tests\nTests manquants",
        4.75,
        2.0,
        2.8,
        2.25,
        WHITE,
        CYAN,
        17,
        15,
    )
    arrow(slide, 7.75, 2.85, 0.7, BLUE)
    card(
        slide,
        "Conclut",
        "Blocker / Important\nMinor / None\nAucun fichier modifié",
        8.65,
        2.0,
        3.45,
        2.25,
        PALE_ORANGE,
        ORANGE,
        17,
        14,
    )
    text(
        slide,
        "Une revue structurée l’attention ; elle ne remplace pas le jugement humain.",
        1.0,
        5.45,
        11.2,
        0.45,
        20,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    notes.append(
        """Durée : 3 min. Utiliser review-change en lecture seule. Comparer demande, plan, rules, diff et résultats. Relever une conformité et demander une conclusion explicite sur les blockers. La revue ne modifie pas les fichiers. Secours : diff et contrôles attendus. Transition : Des frameworks existent si l’on veut plus de structure, mais ils ne sont pas obligatoires."""
    )

    # 14
    slide = new_slide(
        14, "Frameworks existants", "Industrialisation facultative", "3 min"
    )
    card(
        slide,
        "Spec Kit",
        "spec → plan → tasks\n→ implement / converge",
        0.75,
        1.95,
        2.75,
        2.15,
        PALE_BLUE,
        BLUE,
        17,
        15,
    )
    card(
        slide,
        "OpenSpec",
        "explore → propose\n→ apply → archive",
        3.75,
        1.95,
        2.75,
        2.15,
        PALE_CYAN,
        CYAN,
        17,
        15,
    )
    card(
        slide,
        "BMAD Method",
        "Workflows adaptatifs\nspécialisés par rôles",
        6.75,
        1.95,
        2.75,
        2.15,
        PALE_ORANGE,
        ORANGE,
        17,
        15,
    )
    card(
        slide,
        "AIDD",
        "Cycle logiciel plus large\nskills · agents · règles",
        9.75,
        1.95,
        2.75,
        2.15,
        PALE_RED,
        RED,
        17,
        15,
    )
    text(
        slide,
        "Commencer petit : quelques rules + quelques skills. Industrialiser quand le besoin est réel.",
        1.0,
        5.35,
        11.25,
        0.55,
        19,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    notes.append(
        """Durée : 3 min. Présenter comme des niveaux d’industrialisation facultatifs : Spec Kit, OpenSpec, BMAD Method et AIDD. Ne pas en faire un catalogue ni évoquer une incompatibilité avec Kilo Code. Dire qu’ils fournissent souvent prompts, commandes ou skills pour planifier, implémenter et revoir. On peut commencer petit puis industrialiser. Transition : Le dernier niveau n’est pas un outil : ce sont nos limites et notre responsabilité."""
    )

    # 15
    slide = new_slide(
        15, "Sécurité, bonnes pratiques, conclusion", "À retenir", "4 min"
    )
    card(
        slide,
        "5 gestes",
        "1  Contexte pertinent\n2  Rules explicites\n3  Skills réutilisables\n4  Approbation humaine\n5  Harnais déterministe",
        0.8,
        1.85,
        4.25,
        3.2,
        PALE_BLUE,
        BLUE,
        18,
        16,
    )
    card(
        slide,
        "Toujours",
        "Pas de secrets\nPermissions limitées\nDiff relu\nBranche / worktree\nOutils approuvés",
        5.4,
        1.85,
        3.0,
        3.2,
        PALE_ORANGE,
        ORANGE,
        18,
        15,
    )
    card(
        slide,
        "TDD ?",
        "Une boucle rouge / vert / refactor\npeut être lisible pour l’agent.\n\nPiste intéressante, pas obligation\net pas une promesse de maîtrise.",
        8.75,
        1.85,
        3.7,
        3.2,
        PALE_CYAN,
        CYAN,
        18,
        14,
    )
    text(
        slide,
        "Ressources : docs/resources.md  ·  Questions : 10 min",
        1.0,
        5.72,
        11.3,
        0.35,
        16,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    notes.append(
        """Durée : 4 min. Rappeler : aucun secret dans prompts, rules, logs ou démo ; vérifier les commandes ; limiter permissions ; branche/worktree ; diff ; pas de production ni données sensibles ; rules ≠ barrières techniques. Dire avec prudence : « Le TDD semble fournir une boucle particulièrement lisible pour l’agent. C’est une piste intéressante, mais ce n’est ni obligatoire ni une pratique que je prétends aujourd’hui maîtriser avec l’IA. » Conclure : contexte, rules, skills, approbation, harnais. Passer la parole pour 8 minutes de formations, puis 10 minutes de questions. Les ressources officielles sont dans docs/resources.md."""
    )

    if len(notes) != 15:
        raise RuntimeError(f"Expected 15 notes, got {len(notes)}")
    for slide, speaker_notes in zip(prs.slides, notes, strict=True):
        add_notes(slide, speaker_notes)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    make_deck()
